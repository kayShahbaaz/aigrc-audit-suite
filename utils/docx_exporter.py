"""
utils/docx_exporter.py — AI GRC Audit Suite
==============================================
Professional DOCX document generation for Modules 1, 4, 5, and 6.

All DOCX exports in the app go through here.
Modules never import python-docx directly — they call these functions.

Why centralise here:
- Consistent formatting across all documents (fonts, colours, spacing)
- One place to update branding or layout for all outputs
- Modules stay focused on content logic, not document formatting

Functions:
    export_policy(content, company_profile)         → bytes  (Module 1)
    export_gap_assessment(content, company_profile) → bytes  (Module 4)
    export_vendor_risk(content, company_profile)    → bytes  (Module 5)
    export_ir_playbook(content, company_profile)    → bytes  (Module 6)

All functions return bytes so Streamlit's st.download_button() can serve them directly.
"""

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# =============================================================
# BRAND COLOURS — defined once, used throughout
# Professional dark navy + accent gold — works for audit docs
# =============================================================

COLOR_NAVY   = RGBColor(0x1A, 0x37, 0x5E)   # #1A375E — headings
COLOR_GOLD   = RGBColor(0xC8, 0x9A, 0x1A)   # #C89A1A — accent / dividers
COLOR_DARK   = RGBColor(0x1F, 0x1F, 0x1F)   # #1F1F1F — body text
COLOR_GREY   = RGBColor(0x55, 0x55, 0x55)   # #555555 — secondary text
COLOR_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)   # #FFFFFF — table header text
COLOR_RED    = RGBColor(0xC0, 0x39, 0x2B)   # #C0392B — critical/high risk
COLOR_AMBER  = RGBColor(0xE6, 0x7E, 0x22)   # #E67E22 — medium risk
COLOR_GREEN  = RGBColor(0x27, 0xAE, 0x60)   # #27AE60 — low risk / met


# =============================================================
# INTERNAL HELPERS
# Private functions used by the public export functions below
# =============================================================

def _new_document() -> Document:
    """
    Create a new Document with base styles applied.
    Sets default font to Calibri 11pt — professional, widely supported.
    """
    doc = Document()

    # Set page margins — 2.5cm on all sides, standard for business docs
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Set default body font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = COLOR_DARK

    return doc


def _add_cover_page(doc: Document, title: str, subtitle: str, company_profile: dict):
    """
    Add a professional cover page to the document.
    Includes document title, company name, auditor, and date.
    Called at the start of every export function.
    """
    # Vertical spacer — push title down from top of page
    for _ in range(6):
        doc.add_paragraph()

    # Document title — large, bold, navy
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = COLOR_NAVY
    run.font.name = "Calibri"

    # Subtitle — smaller, gold colour
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p_subtitle.add_run(subtitle)
    run2.bold = False
    run2.font.size = Pt(14)
    run2.font.color.rgb = COLOR_GOLD
    run2.font.name = "Calibri"

    doc.add_paragraph()  # spacer

    # Horizontal rule — gold line under title
    _add_horizontal_rule(doc, color=COLOR_GOLD)

    doc.add_paragraph()  # spacer

    # Company and engagement details
    details = [
        ("Prepared for",  company_profile.get("company_name", "Client Organisation")),
        ("Industry",      company_profile.get("industry", "—")),
        ("Prepared by",   company_profile.get("auditor_name", "GRC Auditor")),
        ("Engagement",    company_profile.get("engagement_date", datetime.today().strftime("%d %B %Y"))),
        ("Date",          datetime.today().strftime("%d %B %Y")),
        ("Version",       "1.0"),
        ("Classification","CONFIDENTIAL"),
    ]

    for label, value in details:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = p.add_run(f"{label}:  ")
        label_run.bold = True
        label_run.font.size = Pt(11)
        label_run.font.color.rgb = COLOR_GREY
        value_run = p.add_run(value)
        value_run.bold = False
        value_run.font.size = Pt(11)
        value_run.font.color.rgb = COLOR_DARK

    # Page break — everything after this is the document body
    doc.add_page_break()


def _add_heading(doc: Document, text: str, level: int = 1):
    """
    Add a formatted heading to the document.
    Level 1 = section heading (navy, 14pt, bold)
    Level 2 = subsection heading (navy, 12pt, bold)
    Level 3 = minor heading (dark, 11pt, bold)
    """
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"

    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_NAVY
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after  = Pt(6)
        # Add a thin gold border below level-1 headings
        _add_bottom_border(p, color="C89A1A")

    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_NAVY
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after  = Pt(4)

    else:  # level 3
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_DARK
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(2)


def _add_body_text(doc: Document, text: str):
    """
    Add a paragraph of normal body text.
    Calibri 11pt, dark colour, 6pt spacing after.
    """
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_DARK
    p.paragraph_format.space_after = Pt(6)


def _add_numbered_clause(doc: Document, number: str, text: str):
    """
    Add a numbered clause — used in policy documents.
    e.g. "3.1  All staff must complete annual security awareness training."
    """
    p = doc.add_paragraph()
    # Number in bold gold
    num_run = p.add_run(f"{number}  ")
    num_run.bold = True
    num_run.font.color.rgb = COLOR_GOLD
    num_run.font.size = Pt(11)
    # Clause text in normal dark
    text_run = p.add_run(text)
    text_run.font.size = Pt(11)
    text_run.font.color.rgb = COLOR_DARK
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(4)


def _add_horizontal_rule(doc: Document, color: RGBColor = COLOR_NAVY):
    """
    Add a thin horizontal line — used as a section divider.
    Creates a paragraph with a bottom border which renders as a visible line.
    """
    p = doc.add_paragraph()
    _add_bottom_border(p, color="C89A1A" if color == COLOR_GOLD else "1A375E")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)


def _add_bottom_border(paragraph, color: str = "1A375E", size: int = 6):
    """
    Add a bottom border to a paragraph using OOXML directly.
    python-docx doesn't expose paragraph borders natively so we use XML.
    color: hex string without # (e.g. '1A375E')
    size:  border thickness in eighths of a point (6 = 0.75pt line)
    """
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_info_box(doc: Document, label: str, value: str):
    """
    Add a key-value info line — used in summary sections.
    e.g.  Risk Level:   HIGH
    """
    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}:   ")
    label_run.bold = True
    label_run.font.size = Pt(11)
    label_run.font.color.rgb = COLOR_NAVY
    value_run = p.add_run(value)
    value_run.font.size = Pt(11)
    value_run.font.color.rgb = COLOR_DARK
    p.paragraph_format.space_after = Pt(3)


def _add_table_header_row(table, headers: list):
    """
    Style the first row of a table as a dark navy header row.
    Text is white and bold. Called right after creating a table.
    """
    # The first row is the header row
    header_row = table.rows[0]
    for i, header_text in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = header_text
        # Set cell background to navy
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1A375E")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)
        # Format the text
        run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(header_text)
        run.bold = True
        run.font.color.rgb = COLOR_WHITE
        run.font.size = Pt(10)
        run.font.name = "Calibri"
        cell.paragraphs[0].clear()
        run2 = cell.paragraphs[0].add_run(header_text)
        run2.bold = True
        run2.font.color.rgb = COLOR_WHITE
        run2.font.size = Pt(10)
        run2.font.name = "Calibri"


def _doc_to_bytes(doc: Document) -> bytes:
    """
    Save the Document object to an in-memory bytes buffer and return it.
    Streamlit's download_button needs bytes — this is how we provide them
    without writing a temp file to disk.
    """
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


# =============================================================
# PUBLIC EXPORT FUNCTIONS
# =============================================================

def export_policy(content: dict, company_profile: dict) -> bytes:
    """
    Generate a professional compliance policy document in DOCX format.
    Called by modules/policy.py after the LLM generates the content.

    Expected content dict keys:
        framework_name   : str  — e.g. "NCA ECC"
        policy_type      : str  — e.g. "Access Control Policy"
        scope            : str  — who the policy applies to
        policy_body      : str  — main policy text from LLM (numbered clauses)
        roles            : str  — roles and responsibilities section
        review_schedule  : str  — how often policy is reviewed
        arabic_summary   : str  — optional Arabic translation (bilingual mode)

    Returns:
        bytes — ready for st.download_button()
    """
    doc = _new_document()

    framework = content.get("framework_name", "Compliance Framework")
    policy_type = content.get("policy_type", "Security Policy")
    company = company_profile.get("company_name", "Organisation")

    # Cover page
    _add_cover_page(
        doc,
        title=policy_type,
        subtitle=f"{framework} Compliance Policy | {company}",
        company_profile=company_profile,
    )

    # --- 1. Purpose and Scope ---
    _add_heading(doc, "1. Purpose and Scope", level=1)
    _add_body_text(doc, content.get("scope", ""))

    # --- 2. Policy Statement ---
    _add_heading(doc, "2. Policy Statement", level=1)
    # The policy body comes as numbered clauses from the LLM
    # We render them as-is since the LLM formats them as "2.1  ...", "2.2  ..."
    policy_body = content.get("policy_body", "")
    for line in policy_body.split("\n"):
        line = line.strip()
        if not line:
            continue
        # Detect numbered clause (starts with digit and dot, e.g. "2.1  " or "3.")
        import re
        if re.match(r"^\d+[\.\d]*\s+", line):
            # Split the number from the rest
            parts = re.split(r"(\d+[\.\d]*)\s+", line, maxsplit=1)
            if len(parts) >= 3:
                _add_numbered_clause(doc, parts[1], parts[2])
            else:
                _add_body_text(doc, line)
        else:
            _add_body_text(doc, line)

    # --- 3. Roles and Responsibilities ---
    _add_heading(doc, "3. Roles and Responsibilities", level=1)
    _add_body_text(doc, content.get("roles", ""))

    # --- 4. Review Schedule ---
    _add_heading(doc, "4. Review and Maintenance", level=1)
    review = content.get("review_schedule", "This policy shall be reviewed annually.")
    _add_body_text(doc, review)
    _add_body_text(doc,
        "Any significant changes to the organisation's risk profile, technology environment, "
        "or regulatory requirements may trigger an out-of-cycle review."
    )

    # --- 5. Compliance and Exceptions ---
    _add_heading(doc, "5. Compliance and Exceptions", level=1)
    _add_body_text(doc,
        f"All personnel within scope of this policy are required to comply with its requirements. "
        f"Non-compliance may result in disciplinary action in accordance with {company}'s "
        f"HR policies and applicable laws. Requests for exceptions must be submitted in writing "
        f"to the Information Security function and approved by the CISO or equivalent."
    )

    # --- 6. Approval Section ---
    _add_heading(doc, "6. Document Approval", level=1)
    _add_body_text(doc, "This policy has been reviewed and approved by the following:")

    # Approval table
    approval_table = doc.add_table(rows=4, cols=4)
    approval_table.style = "Table Grid"
    _add_table_header_row(approval_table, ["Role", "Name", "Signature", "Date"])

    roles_for_approval = ["Policy Owner", "CISO / Security Lead", "Executive Sponsor"]
    for i, role in enumerate(roles_for_approval):
        row = approval_table.rows[i + 1]
        row.cells[0].text = role
        row.cells[1].text = ""
        row.cells[2].text = ""
        row.cells[3].text = ""

    doc.add_paragraph()  # spacer after table

    # --- Arabic Summary (bilingual mode) ---
    arabic_summary = content.get("arabic_summary", "")
    if arabic_summary:
        doc.add_page_break()
        _add_heading(doc, "ملخص السياسة — Policy Summary (Arabic)", level=1)
        _add_body_text(doc, "يُرجى الاطلاع على النص الإنجليزي الكامل للسياسة. فيما يلي ملخص تنفيذي باللغة العربية.")
        p = doc.add_paragraph()
        run = p.add_run(arabic_summary)
        run.font.name = "Arial"     # Arial has full Arabic glyph support
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_DARK
        # Right-to-left paragraph direction
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    return _doc_to_bytes(doc)


def export_gap_assessment(content: dict, company_profile: dict) -> bytes:
    """
    Generate a gap assessment report in DOCX format.
    Called by modules/gap.py after RAG analysis is complete.

    Expected content dict keys:
        framework_name     : str   — e.g. "NCA ECC"
        executive_summary  : str   — high-level summary
        overall_score      : float — 0.0 to 1.0 (e.g. 0.62 = 62%)
        controls_met       : list  — list of {id, name, domain}
        controls_partial   : list  — list of {id, name, domain, gap}
        controls_not_met   : list  — list of {id, name, domain, gap}
        priority_actions   : list  — top 10 recommended actions
        arabic_summary     : str   — optional Arabic translation
    """
    doc = _new_document()

    framework = content.get("framework_name", "Framework")
    company = company_profile.get("company_name", "Organisation")

    _add_cover_page(
        doc,
        title="Gap Assessment Report",
        subtitle=f"{framework} | {company}",
        company_profile=company_profile,
    )

    # --- Executive Summary ---
    _add_heading(doc, "1. Executive Summary", level=1)
    _add_body_text(doc, content.get("executive_summary", ""))

    # Compliance score callout
    score = content.get("overall_score", 0.0)
    score_pct = f"{score * 100:.0f}%"
    met_count     = len(content.get("controls_met", []))
    partial_count = len(content.get("controls_partial", []))
    not_met_count = len(content.get("controls_not_met", []))
    total_count   = met_count + partial_count + not_met_count

    doc.add_paragraph()
    _add_info_box(doc, "Overall Compliance Score", score_pct)
    _add_info_box(doc, "Total Controls Assessed",  str(total_count))
    _add_info_box(doc, "Controls Met",             str(met_count))
    _add_info_box(doc, "Controls Partially Met",   str(partial_count))
    _add_info_box(doc, "Controls Not Met",         str(not_met_count))
    doc.add_paragraph()

    # --- Priority Action Plan ---
    _add_heading(doc, "2. Priority Action Plan", level=1)
    _add_body_text(doc, "The following actions are recommended in priority order based on risk impact:")
    doc.add_paragraph()

    for i, action in enumerate(content.get("priority_actions", []), start=1):
        _add_numbered_clause(doc, str(i), action)

    # --- Controls Not Met ---
    _add_heading(doc, "3. Controls Not Met", level=1)
    not_met = content.get("controls_not_met", [])
    if not_met:
        table = doc.add_table(rows=len(not_met) + 1, cols=4)
        table.style = "Table Grid"
        _add_table_header_row(table, ["Control ID", "Control Name", "Domain", "Gap Description"])
        for i, ctrl in enumerate(not_met):
            row = table.rows[i + 1]
            row.cells[0].text = ctrl.get("id", "")
            row.cells[1].text = ctrl.get("name", "")
            row.cells[2].text = ctrl.get("domain", "")
            row.cells[3].text = ctrl.get("gap", "No evidence found")
    else:
        _add_body_text(doc, "No controls were assessed as Not Met.")
    doc.add_paragraph()

    # --- Controls Partially Met ---
    _add_heading(doc, "4. Controls Partially Met", level=1)
    partial = content.get("controls_partial", [])
    if partial:
        table2 = doc.add_table(rows=len(partial) + 1, cols=4)
        table2.style = "Table Grid"
        _add_table_header_row(table2, ["Control ID", "Control Name", "Domain", "Gap Description"])
        for i, ctrl in enumerate(partial):
            row = table2.rows[i + 1]
            row.cells[0].text = ctrl.get("id", "")
            row.cells[1].text = ctrl.get("name", "")
            row.cells[2].text = ctrl.get("domain", "")
            row.cells[3].text = ctrl.get("gap", "Partial evidence found")
    else:
        _add_body_text(doc, "No controls were assessed as Partially Met.")
    doc.add_paragraph()

    # --- Controls Met (summary table) ---
    _add_heading(doc, "5. Controls Met", level=1)
    met = content.get("controls_met", [])
    if met:
        table3 = doc.add_table(rows=len(met) + 1, cols=3)
        table3.style = "Table Grid"
        _add_table_header_row(table3, ["Control ID", "Control Name", "Domain"])
        for i, ctrl in enumerate(met):
            row = table3.rows[i + 1]
            row.cells[0].text = ctrl.get("id", "")
            row.cells[1].text = ctrl.get("name", "")
            row.cells[2].text = ctrl.get("domain", "")
    else:
        _add_body_text(doc, "No controls were assessed as fully Met.")

    # --- Arabic Summary (bilingual mode) ---
    arabic_summary = content.get("arabic_summary", "")
    if arabic_summary:
        doc.add_page_break()
        _add_heading(doc, "ملخص تنفيذي — Executive Summary (Arabic)", level=1)
        p = doc.add_paragraph()
        run = p.add_run(arabic_summary)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_DARK
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    return _doc_to_bytes(doc)


def export_vendor_risk(content: dict, company_profile: dict) -> bytes:
    """
    Generate a vendor risk assessment report in DOCX format.
    Called by modules/vendor.py.

    Expected content dict keys:
        vendor_name        : str   — vendor being assessed
        service_type       : str   — what the vendor provides
        risk_level         : str   — Low / Medium / High / Critical
        risk_score         : int   — 0–100
        red_flags          : list  — list of identified risk flags (strings)
        contract_clauses   : list  — recommended contract clauses
        due_diligence      : list  — due diligence checklist items
        summary            : str   — narrative summary
        arabic_summary     : str   — optional Arabic translation
    """
    doc = _new_document()

    vendor = content.get("vendor_name", "Vendor")
    company = company_profile.get("company_name", "Organisation")

    _add_cover_page(
        doc,
        title="Vendor Risk Assessment",
        subtitle=f"{vendor} | Prepared for {company}",
        company_profile=company_profile,
    )

    # --- Vendor Overview ---
    _add_heading(doc, "1. Vendor Overview", level=1)
    _add_info_box(doc, "Vendor Name",    content.get("vendor_name", "—"))
    _add_info_box(doc, "Service Type",   content.get("service_type", "—"))
    _add_info_box(doc, "Risk Score",     f"{content.get('risk_score', 0)} / 100")
    _add_info_box(doc, "Risk Level",     content.get("risk_level", "—"))
    doc.add_paragraph()
    _add_body_text(doc, content.get("summary", ""))

    # --- Red Flags ---
    _add_heading(doc, "2. Risk Red Flags Identified", level=1)
    red_flags = content.get("red_flags", [])
    if red_flags:
        for i, flag in enumerate(red_flags, start=1):
            _add_numbered_clause(doc, str(i), flag)
    else:
        _add_body_text(doc, "No red flags identified during this assessment.")

    # --- Recommended Contract Clauses ---
    _add_heading(doc, "3. Recommended Contract Clauses", level=1)
    _add_body_text(doc, "The following clauses are recommended for inclusion in the vendor contract or SLA:")
    doc.add_paragraph()
    for i, clause in enumerate(content.get("contract_clauses", []), start=1):
        _add_numbered_clause(doc, str(i), clause)

    # --- Due Diligence Checklist ---
    _add_heading(doc, "4. Due Diligence Checklist", level=1)
    checklist = content.get("due_diligence", [])
    if checklist:
        table = doc.add_table(rows=len(checklist) + 1, cols=3)
        table.style = "Table Grid"
        _add_table_header_row(table, ["#", "Due Diligence Item", "Status"])
        for i, item in enumerate(checklist):
            row = table.rows[i + 1]
            row.cells[0].text = str(i + 1)
            row.cells[1].text = item
            row.cells[2].text = "☐ Pending"

    # --- Arabic Summary ---
    arabic_summary = content.get("arabic_summary", "")
    if arabic_summary:
        doc.add_page_break()
        _add_heading(doc, "ملخص تقييم مخاطر المورد — Vendor Risk Summary (Arabic)", level=1)
        p = doc.add_paragraph()
        run = p.add_run(arabic_summary)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_DARK
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    return _doc_to_bytes(doc)


def export_ir_playbook(content: dict, company_profile: dict) -> bytes:
    """
    Generate an Incident Response Playbook in DOCX format.
    Called by modules/playbook.py.

    Expected content dict keys:
        incident_type          : str   — e.g. "Ransomware"
        classification         : str   — severity and impact classification
        response_team          : list  — {role, responsibilities}
        response_timeline      : list  — {phase, step, action, owner, timeframe}
        escalation_matrix      : list  — {trigger, escalate_to, method}
        comms_templates        : list  — {audience, template_text}
        regulatory_requirements: list  — {body, requirement, deadline}
        lessons_learned        : str   — lessons learned template text
        arabic_summary         : str   — optional Arabic translation
    """
    doc = _new_document()

    incident_type = content.get("incident_type", "Incident")
    company = company_profile.get("company_name", "Organisation")

    _add_cover_page(
        doc,
        title=f"{incident_type} Response Playbook",
        subtitle=f"Incident Response Plan | {company}",
        company_profile=company_profile,
    )

    # --- Classification ---
    _add_heading(doc, "1. Incident Classification", level=1)
    _add_body_text(doc, content.get("classification", ""))

    # --- Response Team ---
    _add_heading(doc, "2. Response Team Roles", level=1)
    team = content.get("response_team", [])
    if team:
        table = doc.add_table(rows=len(team) + 1, cols=2)
        table.style = "Table Grid"
        _add_table_header_row(table, ["Role", "Responsibilities"])
        for i, member in enumerate(team):
            row = table.rows[i + 1]
            row.cells[0].text = member.get("role", "")
            row.cells[1].text = member.get("responsibilities", "")
    doc.add_paragraph()

    # --- Response Timeline ---
    _add_heading(doc, "3. Step-by-Step Response Timeline", level=1)
    timeline = content.get("response_timeline", [])
    if timeline:
        table2 = doc.add_table(rows=len(timeline) + 1, cols=4)
        table2.style = "Table Grid"
        _add_table_header_row(table2, ["Phase", "Action", "Owner", "Timeframe"])
        for i, step in enumerate(timeline):
            row = table2.rows[i + 1]
            row.cells[0].text = step.get("phase", "")
            row.cells[1].text = step.get("action", "")
            row.cells[2].text = step.get("owner", "")
            row.cells[3].text = step.get("timeframe", "")
    doc.add_paragraph()

    # --- Escalation Matrix ---
    _add_heading(doc, "4. Escalation Matrix", level=1)
    escalation = content.get("escalation_matrix", [])
    if escalation:
        table3 = doc.add_table(rows=len(escalation) + 1, cols=3)
        table3.style = "Table Grid"
        _add_table_header_row(table3, ["Trigger Condition", "Escalate To", "Notification Method"])
        for i, item in enumerate(escalation):
            row = table3.rows[i + 1]
            row.cells[0].text = item.get("trigger", "")
            row.cells[1].text = item.get("escalate_to", "")
            row.cells[2].text = item.get("method", "")
    doc.add_paragraph()

    # --- Communication Templates ---
    _add_heading(doc, "5. Communication Templates", level=1)
    for tmpl in content.get("comms_templates", []):
        _add_heading(doc, tmpl.get("audience", ""), level=2)
        _add_body_text(doc, tmpl.get("template_text", ""))

    # --- Regulatory Requirements ---
    _add_heading(doc, "6. Regulatory Notification Requirements", level=1)
    reg_reqs = content.get("regulatory_requirements", [])
    if reg_reqs:
        table4 = doc.add_table(rows=len(reg_reqs) + 1, cols=3)
        table4.style = "Table Grid"
        _add_table_header_row(table4, ["Regulatory Body", "Requirement", "Deadline"])
        for i, req in enumerate(reg_reqs):
            row = table4.rows[i + 1]
            row.cells[0].text = req.get("body", "")
            row.cells[1].text = req.get("requirement", "")
            row.cells[2].text = req.get("deadline", "")
    doc.add_paragraph()

    # --- Lessons Learned ---
    _add_heading(doc, "7. Lessons Learned Template", level=1)
    _add_body_text(doc, content.get("lessons_learned", ""))

    # --- Arabic Summary ---
    arabic_summary = content.get("arabic_summary", "")
    if arabic_summary:
        doc.add_page_break()
        _add_heading(doc, "ملخص دليل الاستجابة للحوادث — IR Playbook Summary (Arabic)", level=1)
        p = doc.add_paragraph()
        run = p.add_run(arabic_summary)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_DARK
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    return _doc_to_bytes(doc)
